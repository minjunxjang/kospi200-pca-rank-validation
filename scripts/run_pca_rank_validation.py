from __future__ import annotations

import json
import math
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ROOT = PROJECT_ROOT / "data" / "raw"
OUT = PROJECT_ROOT / "results" / "recomputed" / "intermediate"
OUT.mkdir(parents=True, exist_ok=True)
CSV_OUT = PROJECT_ROOT / "results" / "recomputed" / "csv"
CSV_OUT.mkdir(parents=True, exist_ok=True)

T_EXPECTED = 4582
N_EXPECTED = 407
N_VALUES = [5, 10, 20, 60]
K_VALUES = [5, 10, 20, 60]
W_VALUES = [60, 120, 252]
PC_VALUES = ["PC1", "PC12", "PC123"]
MIN_STOCKS = 50


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).split()[0]
    if "-" in s:
        return datetime.strptime(s, "%Y-%m-%d").date()
    return datetime.strptime(s, "%Y%m%d").date()


def read_wide_sheet(wb, sheet_name, expected_t=T_EXPECTED, expected_n=N_EXPECTED):
    ws = wb[sheet_name]
    rows = ws.iter_rows(min_row=1, max_row=expected_t + 1, min_col=1, max_col=expected_n + 1, values_only=True)
    header = next(rows)
    tickers = [str(x).zfill(6) if x is not None else "" for x in header[1:]]
    dates = np.empty(expected_t, dtype="datetime64[D]")
    data = np.full((expected_t, expected_n), np.nan, dtype=np.float64)
    for i, row in enumerate(rows):
        if i >= expected_t:
            break
        dates[i] = np.datetime64(as_date(row[0]))
        vals = row[1:]
        data[i, :] = [np.nan if v is None else float(v) for v in vals]
    if i + 1 != expected_t:
        raise ValueError(f"{sheet_name}: expected {expected_t} rows, got {i + 1}")
    return dates, tickers, data


def rolling_sum_strict(x, window):
    finite = np.isfinite(x)
    values = np.where(finite, x, 0.0)
    cs = np.vstack([np.zeros((1, x.shape[1])), np.cumsum(values, axis=0)])
    cc = np.vstack([np.zeros((1, x.shape[1]), dtype=np.int32), np.cumsum(finite, axis=0, dtype=np.int32)])
    out = np.full_like(x, np.nan)
    sums = cs[window:] - cs[:-window]
    counts = cc[window:] - cc[:-window]
    out[window - 1:] = np.where(counts == window, sums, np.nan)
    return out


def derive_cum20_from_cum10(cum10):
    out = np.full_like(cum10, np.nan)
    out[10:] = cum10[10:] + cum10[:-10]
    return out


def row_rank(x):
    return pd.DataFrame(x).rank(axis=1, method="average", na_option="keep").to_numpy(dtype=np.float64)


def daily_rank_metrics(signal, future):
    valid = np.isfinite(signal) & np.isfinite(future)
    n = valid.sum(axis=1).astype(np.int32)
    x = np.where(valid, signal, np.nan)
    y = np.where(valid, future, np.nan)
    xr = row_rank(x)
    yr = row_rank(y)
    xm = np.nanmean(xr, axis=1)
    ym = np.nanmean(yr, axis=1)
    xc = xr - xm[:, None]
    yc = yr - ym[:, None]
    num = np.nansum(xc * yc, axis=1)
    den = np.sqrt(np.nansum(xc * xc, axis=1) * np.nansum(yc * yc, axis=1))
    ic = np.divide(num, den, out=np.full(signal.shape[0], np.nan), where=den > 0)
    denom = np.maximum(n[:, None], 1)
    xq = np.ceil(5.0 * xr / denom)
    yq = np.ceil(5.0 * yr / denom)
    xq = np.clip(xq, 1, 5)
    yq = np.clip(yq, 1, 5)
    xqm = np.nanmean(xq, axis=1)
    yqm = np.nanmean(yq, axis=1)
    xqc = xq - xqm[:, None]
    yqc = yq - yqm[:, None]
    qnum = np.nansum(xqc * yqc, axis=1)
    qden = np.sqrt(np.nansum(xqc * xqc, axis=1) * np.nansum(yqc * yqc, axis=1))
    quintile_ic = np.divide(qnum, qden, out=np.full(signal.shape[0], np.nan), where=qden > 0)
    ic[n < MIN_STOCKS] = np.nan
    quintile_ic[n < MIN_STOCKS] = np.nan
    xq[n < MIN_STOCKS, :] = np.nan
    yq[n < MIN_STOCKS, :] = np.nan
    return ic, quintile_ic, n, xq, yq


def normal_cdf(x):
    return 0.5 * (1.0 + math.erf(x / math.sqrt(2.0)))


def hac_mean_test(values, null, lag):
    x = np.asarray(values, dtype=float)
    x = x[np.isfinite(x)]
    n = x.size
    if n < 3:
        return dict(obs=n, mean=np.nan, excess=np.nan, se=np.nan, t=np.nan, p_two=np.nan, p_pos=np.nan, lag=0)
    mean = float(x.mean())
    u = x - mean
    lag = int(min(max(lag, 0), n - 1))
    lrv = float(np.dot(u, u) / n)
    for ell in range(1, lag + 1):
        gamma = float(np.dot(u[ell:], u[:-ell]) / n)
        lrv += 2.0 * (1.0 - ell / (lag + 1.0)) * gamma
    se = math.sqrt(max(lrv, 0.0) / n)
    t = (mean - null) / se if se > 0 else np.nan
    p_two = math.erfc(abs(t) / math.sqrt(2.0)) if np.isfinite(t) else np.nan
    p_pos = 0.5 * math.erfc(t / math.sqrt(2.0)) if np.isfinite(t) else np.nan
    return dict(obs=n, mean=mean, excess=mean - null, se=se, t=t, p_two=p_two, p_pos=p_pos, lag=lag)


def bh_qvalues(p):
    p = np.asarray(p, dtype=float)
    q = np.full_like(p, np.nan)
    valid_idx = np.flatnonzero(np.isfinite(p))
    if valid_idx.size == 0:
        return q
    pv = p[valid_idx]
    order = np.argsort(pv)
    ranked = pv[order]
    m = ranked.size
    adj = ranked * m / np.arange(1, m + 1)
    adj = np.minimum.accumulate(adj[::-1])[::-1]
    adj = np.minimum(adj, 1.0)
    back = np.empty_like(adj)
    back[order] = adj
    q[valid_idx] = back
    return q


def case_table():
    rows = []
    for n in N_VALUES:
        rows.append(dict(case_id=f"RAW_N{n}", W="RAW", pc="RAW", n=n, is_raw=True))
    for w in W_VALUES:
        for pc in PC_VALUES:
            for n in N_VALUES:
                rows.append(dict(case_id=f"W{w}_{pc}_N{n}", W=w, pc=pc, n=n, is_raw=False))
    return pd.DataFrame(rows)


def append_confusion(target, case, k, year, pred, realized):
    mask = np.isfinite(pred) & np.isfinite(realized)
    if not mask.any():
        return
    code = (pred[mask].astype(int) - 1) * 5 + (realized[mask].astype(int) - 1)
    counts = np.bincount(code, minlength=25).reshape(5, 5)
    total = int(counts.sum())
    for i in range(5):
        row_total = int(counts[i].sum())
        for j in range(5):
            target.append({
                "case_id": case,
                "k": k,
                "year": year,
                "signal_quintile": i + 1,
                "future_quintile": j + 1,
                "count": int(counts[i, j]),
                "row_rate": counts[i, j] / row_total if row_total else np.nan,
                "overall_rate": counts[i, j] / total if total else np.nan,
            })


def main():
    cases = case_table()
    c_count = len(cases)
    print(f"Cases: {c_count}", flush=True)

    print("Loading future returns and universe...", flush=True)
    fwb = load_workbook(ROOT / "FUTURE_LOG_RETURNS_5_10_20_60.xlsx", read_only=True, data_only=True, keep_links=False)
    dates, tickers, eligible_raw = read_wide_sheet(fwb, "COMMON252_ELIGIBLE")
    eligible = eligible_raw == 1
    future = {}
    for k in K_VALUES:
        d2, t2, arr = read_wide_sheet(fwb, f"FWD_LOG_{k}")
        if not np.array_equal(dates, d2) or tickers != t2:
            raise ValueError(f"FWD {k} alignment mismatch")
        future[k] = np.where(eligible, arr, np.nan)
        print(f" loaded FWD {k}", flush=True)
    fwb.close()

    years = pd.to_datetime(dates).year.to_numpy()
    year_values = np.unique(years)

    print("Loading raw daily returns and forming n-day signals...", flush=True)
    bwb = load_workbook(ROOT / "00_COMMON252_BASE.xlsx", read_only=True, data_only=True, keep_links=False)
    d2, t2, raw_daily = read_wide_sheet(bwb, "DAILY_LOG_RETURN")
    bwb.close()
    if not np.array_equal(dates, d2) or tickers != t2:
        raise ValueError("Raw alignment mismatch")
    raw_cum = {n: rolling_sum_strict(raw_daily, n) for n in N_VALUES}

    ic_cube = np.full((len(dates), c_count, len(K_VALUES)), np.nan, dtype=np.float32)
    quintile_ic_cube = np.full_like(ic_cube, np.nan)
    n_cube = np.zeros((len(dates), c_count, len(K_VALUES)), dtype=np.int16)
    confusion = []

    signals = {}
    for i, row in cases.iterrows():
        if row.is_raw:
            signals[row.case_id] = -raw_cum[int(row.n)]

    cumulative_files = {
        60: ROOT / "02_PCA_W60_CUMULATIVE.xlsx",
        120: ROOT / "04_PCA_W120_CUMULATIVE.xlsx",
        252: ROOT / "06_PCA_W252_CUMULATIVE.xlsx",
    }

    for w, path in cumulative_files.items():
        print(f"Loading W{w} cumulative signals...", flush=True)
        wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        for pc in PC_VALUES:
            cache = {}
            for n in [5, 10, 60]:
                d3, t3, arr = read_wide_sheet(wb, f"{pc}_CUM{n}")
                if not np.array_equal(dates, d3) or tickers != t3:
                    raise ValueError(f"W{w} {pc} N{n} alignment mismatch")
                cache[n] = arr
            cache[20] = derive_cum20_from_cum10(cache[10])
            for n in N_VALUES:
                signals[f"W{w}_{pc}_N{n}"] = -cache[n]
            print(f" loaded W{w} {pc}", flush=True)
        wb.close()

    if set(signals) != set(cases.case_id):
        raise ValueError("Signal case set mismatch")

    for ci, row in cases.iterrows():
        case = row.case_id
        signal = signals.pop(case)
        print(f"Computing {ci + 1:02d}/{c_count}: {case}", flush=True)
        for ki, k in enumerate(K_VALUES):
            ic, quintile_ic, nobs, pred_q, realized_q = daily_rank_metrics(signal, future[k])
            ic_cube[:, ci, ki] = ic.astype(np.float32)
            quintile_ic_cube[:, ci, ki] = quintile_ic.astype(np.float32)
            n_cube[:, ci, ki] = np.minimum(nobs, np.iinfo(np.int16).max).astype(np.int16)
            append_confusion(confusion, case, k, "FULL", pred_q, realized_q)
            for yr in year_values:
                year_mask = years == yr
                append_confusion(confusion, case, k, int(yr), pred_q[year_mask], realized_q[year_mask])

    raw_index = {int(cases.iloc[i].n): i for i in range(4)}
    delta_ic = np.full_like(ic_cube, np.nan)
    delta_quintile_ic = np.full_like(quintile_ic_cube, np.nan)
    for ci, row in cases.iterrows():
        if not row.is_raw:
            ri = raw_index[int(row.n)]
            delta_ic[:, ci, :] = ic_cube[:, ci, :] - ic_cube[:, ri, :]
            delta_quintile_ic[:, ci, :] = quintile_ic_cube[:, ci, :] - quintile_ic_cube[:, ri, :]

    metric_defs = [
        ("rank_ic", ic_cube, 0.0),
        ("quintile_ic", quintile_ic_cube, 0.0),
        ("delta_rank_ic", delta_ic, 0.0),
        ("delta_quintile_ic", delta_quintile_ic, 0.0),
    ]
    full_rows = []
    annual_rows = []
    for metric, cube, null in metric_defs:
        for ci, row in cases.iterrows():
            if metric.startswith("delta_") and row.is_raw:
                continue
            for ki, k in enumerate(K_VALUES):
                lag = max(int(row.n), k) - 1
                stats = hac_mean_test(cube[:, ci, ki], null, lag)
                full_rows.append({"metric": metric, **row.to_dict(), "k": k, "null": null, **stats})
                for yr in year_values:
                    stats_y = hac_mean_test(cube[years == yr, ci, ki], null, lag)
                    annual_rows.append({"metric": metric, **row.to_dict(), "k": k, "year": int(yr), "null": null, **stats_y})

    full = pd.DataFrame(full_rows)
    annual = pd.DataFrame(annual_rows)
    for metric in full.metric.unique():
        mask = full.metric == metric
        full.loc[mask, "q_pos"] = bh_qvalues(full.loc[mask, "p_pos"].to_numpy())
        full.loc[mask, "q_two"] = bh_qvalues(full.loc[mask, "p_two"].to_numpy())
    for metric in annual.metric.unique():
        mask_m = annual.metric == metric
        annual.loc[mask_m, "q_pos_global_space"] = bh_qvalues(annual.loc[mask_m, "p_pos"].to_numpy())
        annual.loc[mask_m, "q_two_global_space"] = bh_qvalues(annual.loc[mask_m, "p_two"].to_numpy())
        for yr in year_values:
            mask = mask_m & (annual.year == yr)
            annual.loc[mask, "q_pos_within_year"] = bh_qvalues(annual.loc[mask, "p_pos"].to_numpy())
            annual.loc[mask, "q_two_within_year"] = bh_qvalues(annual.loc[mask, "p_two"].to_numpy())

    confusion_df = pd.DataFrame(confusion)
    cases.to_csv(CSV_OUT / "cases.csv", index=False, encoding="utf-8-sig")
    full.to_csv(CSV_OUT / "full_summary.csv", index=False, encoding="utf-8-sig")
    annual.to_csv(CSV_OUT / "annual_summary.csv", index=False, encoding="utf-8-sig")
    confusion_df.to_csv(OUT / "quintile_confusion.csv", index=False, encoding="utf-8-sig")
    daily_long = []
    for ci, row in cases.iterrows():
        for ki, k in enumerate(K_VALUES):
            frame = pd.DataFrame({
                "date": pd.to_datetime(dates),
                "year": years,
                "case_id": row.case_id,
                "k": k,
                "rank_ic": ic_cube[:, ci, ki],
                "quintile_ic": quintile_ic_cube[:, ci, ki],
                "delta_rank_ic": delta_ic[:, ci, ki],
                "delta_quintile_ic": delta_quintile_ic[:, ci, ki],
                "cross_section_n": n_cube[:, ci, ki],
            })
            daily_long.append(frame)
    pd.concat(daily_long, ignore_index=True).to_csv(OUT / "daily_metrics.csv.gz", index=False, compression="gzip")

    np.savez_compressed(
        OUT / "rank_cubes.npz",
        dates=dates.astype("datetime64[D]").astype(str),
        years=year_values,
        case_ids=cases.case_id.to_numpy(dtype=str),
        W=cases.W.astype(str).to_numpy(),
        pc=cases.pc.to_numpy(dtype=str),
        n=cases.n.to_numpy(dtype=np.int16),
        k=np.array(K_VALUES, dtype=np.int16),
        daily_rank_ic=ic_cube,
        daily_quintile_ic=quintile_ic_cube,
        daily_delta_rank_ic=delta_ic,
        daily_delta_quintile_ic=delta_quintile_ic,
        daily_cross_section_n=n_cube,
    )
    metadata = {
        "date_start": str(dates[0]),
        "date_end": str(dates[-1]),
        "years": [int(x) for x in year_values],
        "case_count": int(c_count),
        "combination_count": int(c_count * len(K_VALUES)),
        "delta_combination_count": int((c_count - 4) * len(K_VALUES)),
        "min_cross_section": MIN_STOCKS,
        "n20_derivation": "CUM20(T)=CUM10(T)+CUM10(T-10 trading rows); equivalent to the source rolling-sum definition.",
        "signal": "negative cumulative raw/residual log return",
        "rank_ic": "daily cross-sectional Spearman rank correlation on pairwise-valid stocks",
        "quintile_ic": "daily cross-sectional Pearson correlation between signal quintile number (1-5) and future-return quintile number (1-5); null=0",
        "hac_lag": "max(n,k)-1 with Bartlett weights",
        "multiple_testing": "BH separately by metric for full sample; annual results include within-year and global-space BH q-values",
    }
    (OUT / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Analysis complete", flush=True)


if __name__ == "__main__":
    main()
